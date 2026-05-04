from pptx import Presentation
import os

pptx_path = r'G:\Meine Ablage\BA\AutoProject\Kolloquim.pptx'
output_path = r'C:\Users\user\.gemini\antigravity\scratch\kolloquium_summary.txt'

try:
    prs = Presentation(pptx_path)
    text = ""
    for i, slide in enumerate(prs.slides):
        text += f"--- Slide {i+1} ---\n"
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(text)
    print(f"Successfully extracted text to {output_path}")
except Exception as e:
    print(f"Error: {e}")
